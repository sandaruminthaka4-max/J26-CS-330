import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()

# Margins
for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

# Title
p_title = doc.add_paragraph()
r_title = p_title.add_run("RESEARCH PROJECT LOG BOOK")
r_title.bold = True
r_title.font.size = Pt(16)
r_title.font.color.rgb = RGBColor(15, 23, 42)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Metadata
meta = [
    ("Module", "IT4010 Research Project"),
    ("Project ID", "J26-CS-330"),
    ("Overall Title", "Lightweight SDN Intrusion Detection and Zero-Touch Mitigation System for IoMT Networks"),
    ("Component Title", "DoS Detection via SDN Flow Analysis"),
    ("Student Name", "Maneth M. G. B. M"),
    ("Registration No", "IT23326354"),
    ("Supervisor", "Mr. Kanishka Yapa")
]

p_meta = doc.add_paragraph()
for k, v in meta:
    r_k = p_meta.add_run(f"{k}: ")
    r_k.bold = True
    r_k.font.size = Pt(10)
    r_v = p_meta.add_run(f"{v}\n")
    r_v.font.size = Pt(10)

doc.add_paragraph()

# Table
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

hdr_cells = table.rows[0].cells
hdr_cells[0].width = Inches(1.3)
hdr_cells[1].width = Inches(5.5)

hdr_cells[0].paragraphs[0].add_run("Date").bold = True
hdr_cells[1].paragraphs[0].add_run("Task / Research Activity Details").bold = True

for cell in hdr_cells:
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0284C7"/>')
    cell._tc.get_or_add_tcPr().append(shd)
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(10)

entries = [
    ("02/09/2026", "Initial research discussion with supervisor Mr. Kanishka Yapa regarding component scope. Agreed on focusing strictly on lightweight, flow-based DoS detection for IoMT networks without payload inspection."),
    ("03/09/2026", "Studied the official CICIoMT2024 dataset research paper from the Canadian Institute for Cybersecurity (UNB). Confirmed suitability of physical medical device traffic profiles (ECG monitors, infusion pumps) and single-source DoS attack vectors."),
    ("05/09/2026", "Formulated research gap: existing deep learning SDN IDS induce 80-200ms latency, causing control plane bottlenecks; identified need for sub-50ms lightweight tree-based classification."),
    ("09/09/2026", "Outlined individual component boundaries against group members. Confirmed my component handles OpenFlow feature extraction, ML DoS classification, and REST API alerting, while teammate (Wijerathna) handles SDN flow mitigation/drop rules."),
    ("11/09/2026", "Drafted component research objectives and performance targets: Accuracy > 95%, F1-score > 95%, inference latency < 50ms, model disk footprint < 10MB."),
    ("14/09/2026", "Submitted official Project Proposal Report to supervisor Mr. Kanishka Yapa. Proposal successfully signed and approved."),
    ("17/09/2026", "Conducted detailed literature review on OpenFlow 1.3 flow table statistics (OFPFlowStatsReply). Selected 45 statistical flow-level metrics (packet rates, byte counters, durations, TCP flags) that preserve patient privacy (HIPAA/GDPR compliance)."),
    ("18/09/2026", "Meeting with Mitigation Component teammate (Wijerathna H.M.S.H). Agreed on integration contract: REST API POST /predict interface, input flow JSON structure, and output alert schema (is_dos, confidence, latency_ms)."),
    ("23/09/2026", "Acquired and uncompressed the raw CICIoMT2024 dataset archive (train_iomt.csv & test_iomt.csv, size > 1.7 GB). Set up local directory structure: data/, src/, models/, app/, tests/."),
    ("26/09/2026", "Developed chunk-based streaming scanner (src/preprocessing/scanner.py) to analyze the full 7,160,831 records without memory overflow. Discovered label distribution: 192,732 benign and multiple DoS flood variants (UDP, ICMP, TCP, SYN, MQTT floods)."),
    ("30/09/2026", "Engineered data cleaning pipeline (src/preprocessing/cleaner.py). Stripped non-OpenFlow payload columns, replaced infinite values (inf, -inf) with column maximums, and imputed NaN values."),
    ("02/10/2026", "Implemented binary target encoding: mapped benign to class 0 (Normal) and volumetric attack variants (UDP Flood, ICMP Flood, SYN Flood, MQTT Flood) to class 1 (DoS)."),
    ("05/10/2026", "Applied stratified sampling to construct clean, balanced datasets preserving real IoMT dynamics: exported train.csv (109,000 rows), val.csv (23,400 rows), and test.csv (23,400 rows). Fitted Min-Max preprocessor (models/preprocessor.joblib)."),
    ("08/10/2026", "Implemented and evaluated baseline classifiers: Traditional Static Rate Threshold model and Decision Tree classifier to establish comparative performance benchmark."),
    ("10/10/2026", "Trained ensemble Random Forest classifier (100 trees). Achieved 99.64% accuracy, but model size was relatively heavy (18 MB) with 3.8ms latency."),
    ("12/10/2026", "Developed optimized XGBoost classifier (src/training/train.py) utilizing gradient-boosted trees tuned for low depth and fast inference. Achieved 99.72% Accuracy, 99.86% F1-score, and 0.73% False Positive Rate (FPR) on the test split."),
    ("13/10/2026", "Benchmarked inference latency across 10,000 test iterations. Measured median execution speed of 0.996 ms (50x faster than the 50ms QoS requirement). Serialized model artifact to models/xgboost_model.joblib (compact size: 226 KB)."),
    ("15/10/2026", "Built production FastAPI REST microservice (app/main.py, app/predictor.py, app/schemas.py). Implemented asynchronous request handling, schema validation, timing middleware, and fail-safe defaults (fail-open to Normal)."),
    ("17/10/2026", "Created automated testing suite with pytest (tests/test_api.py, tests/test_features.py, tests/test_model.py, tests/test_preprocessing.py). All 18 unit tests passed successfully."),
    ("19/10/2026", "Developed command-line verification demo (scripts/demo_predict.py) verifying real benign patient telemetry flow vs high-rate DoS attack flow. Verified sub-millisecond execution."),
    ("21/10/2026", "Designed interactive Security Operations Center (SOC) web dashboard (app/dashboard.html). Embedded hospital ward device topology (Infusion Pump, ECG, Ventilator, Gateway) and real-time audio alert system."),
    ("23/10/2026", "Enhanced dashboard with 4-factor DoS Forensic Proof Engine (Rate Anomaly, IAT Collapse, Flow Asymmetry, Ensemble Proof) and Explainable AI (XAI) feature attribution bars to mathematically prove DoS attacks during project defense."),
    ("24/10/2026", "Implemented strict accuracy gate (>= 95% confidence threshold) in the Active Detection Verdict card to prevent false alarms on critical medical devices and display dynamic verification proof badges."),
    ("25/10/2026", "Compiled formal Research Evaluation Report and Microsoft Word submission artifact (docs/Research_Evaluation_Report_J26_CS_330.docx). Verified complete alignment with proposal objectives and WBS milestones.")
]

for idx, (dt, task) in enumerate(entries):
    row = table.add_row()
    c0 = row.cells[0]
    c1 = row.cells[1]
    c0.width = Inches(1.3)
    c1.width = Inches(5.5)
    
    r0 = c0.paragraphs[0].add_run(dt)
    r0.bold = True
    r0.font.size = Pt(9.5)
    
    r1 = c1.paragraphs[0].add_run(task)
    r1.font.size = Pt(9.5)
    
    bg_color = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for c in (c0, c1):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        c._tc.get_or_add_tcPr().append(shd)

output_path = "docs/Research_Log_Book_IT23326354.docx"
doc.save(output_path)
print(f"SUCCESS: Log book docx created at {output_path}")
